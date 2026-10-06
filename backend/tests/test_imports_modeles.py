"""Modèles Excel/CSV et imports des fournisseurs, verres, bons de réception et utilisateurs."""

import io

import pytest
from django.contrib.auth.models import Group
from openpyxl import load_workbook

from apps.achats.models import BonReception, Fournisseur
from apps.securite.models import Affectation, Utilisateur
from apps.stock.models import Article, stock_disponible

from .test_imports import fichier, importer


def test_modeles_excel_et_csv(affecter, client_de, tunis):
    api = client_de(affecter("vendeur", "reseau.view_magasin", portee="magasin", magasin=tunis))
    reponse = api.get("/api/v1/imports/modeles/fournisseurs.xlsx")
    assert reponse.status_code == 200
    assert 'filename="modele-fournisseurs.xlsx"' in reponse["Content-Disposition"]
    classeur = load_workbook(io.BytesIO(reponse.content))
    assert classeur.sheetnames == ["À remplir", "Exemple", "Aide"]
    entetes = [c.value for c in classeur["À remplir"][1]]
    assert entetes[:2] == ["code", "raison_sociale"]
    assert classeur["À remplir"].max_row == 1  # Rien à effacer avant de remplir.
    assert classeur["Exemple"]["B2"].value == "Essilor Tunisie"
    aide = {r[0]: r[1] for r in classeur["Aide"].iter_rows(min_row=2, values_only=True)}
    assert aide["raison_sociale"] == "oui"

    csv = api.get("/api/v1/imports/modeles/utilisateurs.csv")
    assert csv.content.decode("utf-8-sig").startswith("identifiant;prenom;nom;email;mot_de_passe")
    assert api.get("/api/v1/imports/modeles/inconnu.csv").status_code == 404


@pytest.fixture
def acheteur(affecter, client_de, tunis):
    return client_de(
        affecter(
            "acheteur",
            "achats.add_fournisseur",
            "achats.change_fournisseur",
            "achats.add_bonreception",
            "stock.add_article",
            "stock.change_article",
            "stock.add_prixarticle",
            "stock.change_prixarticle",
            portee="magasin",
            magasin=tunis,
        )
    )


def test_import_des_fournisseurs_cree_et_signale_les_existants(acheteur, tunis):
    existant = Fournisseur.objects.create(
        nom="Essilor Tunisie", pays=tunis.pays, telephone="71000000"
    )
    contenu = fichier(
        "raison_sociale;matricule_fiscal;forme_juridique;fodec;regime_tva;ville;telephone\n"
        "essilor tunisie;1111111/A/M/000;SARL;oui;Payer TVA;Tunis;\n"
        "Hoya Lens Tunisie;2222222/B/M/000;SUARL;non;Export;Sfax;74000000\n"
    )
    apercu = importer(acheteur, contenu, "fournisseurs", apercu=True).json()
    assert (apercu["crees"], apercu["modifies"]) == (1, 1)
    assert "Essilor Tunisie existe déjà" in apercu["alertes"][0]["message"]
    assert not Fournisseur.objects.filter(nom="Hoya Lens Tunisie").exists()

    contenu.seek(0)
    assert importer(acheteur, contenu, "fournisseurs").status_code == 200
    existant.refresh_from_db()
    assert (existant.matricule_fiscal, existant.fodec) == ("1111111/A/M/000", True)
    assert existant.telephone == "71000000"  # Case vide : valeur gardée.
    hoya = Fournisseur.objects.get(nom="Hoya Lens Tunisie")
    assert (hoya.forme_juridique, hoya.regime_tva, hoya.code) == (
        "SUARL",
        "export",
        existant.code + 1,
    )


def test_un_matricule_fiscal_ne_sert_qu_une_fois(acheteur, tunis):
    Fournisseur.objects.create(nom="Essilor Tunisie", pays=tunis.pays, matricule_fiscal="111/A")
    reponse = importer(
        acheteur,
        fichier("code;raison_sociale;matricule_fiscal\n;Zeiss;222/B\n;Hoya;222/B\n"),
        "fournisseurs",
    )
    assert reponse.status_code == 400
    assert reponse.json()["erreurs"] == [
        {"ligne": 3, "message": "matricule_fiscal : 222/B déjà en ligne 2."}
    ]
    assert not Fournisseur.objects.filter(nom="Zeiss").exists()


def test_import_des_verres_par_code_fournisseur(acheteur, tunis):
    essilor = Fournisseur.objects.create(nom="Essilor Tunisie", pays=tunis.pays)
    contenu = fichier(
        "reference;libelle;fournisseur;geometrie;indice;prix_ttc;tva\n"
        f"VER-9;Verre unifocal 1.5;{essilor.code};Unifocal;1,5;60;7\n"
    )
    assert importer(acheteur, contenu, "verres").status_code == 200
    verre = Article.objects.get(reference="VER-9")
    assert (verre.famille, verre.fournisseur, verre.sur_commande) == ("verre", essilor, True)
    assert verre.verre.geometrie == "unifocal"

    autre = fichier(
        "reference;libelle;fournisseur;famille\nMON-9;Monture;Essilor Tunisie;Monture\n"
    )
    reponse = importer(acheteur, autre, "verres")
    assert reponse.json()["erreurs"][0]["message"] == (
        "famille : ce fichier n'importe que des verres."
    )


def test_import_des_bons_de_reception(acheteur, tunis, monture):
    lux = Fournisseur.objects.create(nom="Luxottica Tunisie", pays=tunis.pays)
    avant = stock_disponible(tunis, monture)
    contenu = fichier(
        "numero_bl;date_bl;fournisseur;reference;quantite;prix_achat_ht;taux_remise\n"
        f"BL-1;04/10/2026;{lux.code};MON-T;2;100;10\n"
        "BL-1;04/10/2026;Luxottica Tunisie;MON-T;1;100;0\n"
        "BL-2;05/10/2026;Luxottica Tunisie;MON-T;1;90;\n"
    )
    apercu = importer(acheteur, contenu, "receptions", magasin=str(tunis.public_id), apercu=True)
    assert apercu.json()["crees"] == 2
    assert not BonReception.tous.exists()

    contenu.seek(0)
    assert (
        importer(acheteur, contenu, "receptions", magasin=str(tunis.public_id)).status_code == 200
    )
    bl1 = BonReception.tous.get(numero_bl="BL-1")
    assert (bl1.lignes.count(), str(bl1.date_bl), bl1.total_net_ht) == (2, "2026-10-04", 280)
    assert stock_disponible(tunis, monture) == avant + 4

    # Le même BL ne s'importe pas deux fois.
    contenu.seek(0)
    reponse = importer(acheteur, contenu, "receptions", magasin=str(tunis.public_id))
    assert (
        "BL-1 de Luxottica Tunisie est déjà enregistré" in reponse.json()["erreurs"][0]["message"]
    )


@pytest.fixture
def admin_magasin(client_de, tunis):
    admin = Utilisateur.objects.create_user("admin")
    Affectation.objects.create(
        utilisateur=admin,
        role=Group.objects.get(name="Administrateur"),
        portee="magasin",
        magasin=tunis,
    )
    return client_de(admin)


def test_import_des_utilisateurs_avec_plusieurs_profils(admin_magasin, tunis):
    contenu = fichier(
        "identifiant;prenom;nom;mot_de_passe;profil;magasin\n"
        "sarra;Sarra;Jlassi;Provisoire-2026!;Vendeur;T01\n"
        "sarra;;;Provisoire-2026!;caissier;T01\n"
    )
    reponse = importer(admin_magasin, contenu, "utilisateurs")
    assert reponse.status_code == 200, reponse.json()
    sarra = Utilisateur.objects.get(username="sarra")
    assert sarra.check_password("Provisoire-2026!")
    assert sorted(a.role.name for a in sarra.affectations.all()) == ["Caissier", "Vendeur"]

    contenu.seek(0)
    reponse = importer(admin_magasin, contenu, "utilisateurs")
    assert reponse.json()["erreurs"] == [
        {"ligne": 2, "message": "sarra existe déjà : le modifier dans Accès et sécurité."}
    ]


def test_import_des_utilisateurs_inactifs(admin_magasin, tunis):
    contenu = fichier(
        "identifiant;prenom;nom;mot_de_passe;profil;magasin;actif\n"
        "sarra;Sarra;Jlassi;Provisoire-2026!;Vendeur;T01;oui\n"
        "olfa;Olfa;Chartaoui;Provisoire-2026!;Vendeur;T01;Non\n"
        "rim;Rim;Ayari;Provisoire-2026!;Vendeur;T01;peut-être\n"
    )
    erreurs = importer(admin_magasin, contenu, "utilisateurs").json()["erreurs"]
    assert erreurs == [{"ligne": 4, "message": "actif : « peut-être » ; écrire oui ou non."}]

    contenu = fichier(
        "identifiant;prenom;nom;mot_de_passe;profil;magasin;actif\n"
        "sarra;Sarra;Jlassi;Provisoire-2026!;Vendeur;T01;oui\n"
        "olfa;Olfa;Chartaoui;Provisoire-2026!;Vendeur;T01;Non\n"
    )
    assert importer(admin_magasin, contenu, "utilisateurs").status_code == 200
    assert Utilisateur.objects.get(username="sarra").is_active
    assert not Utilisateur.objects.get(username="olfa").is_active


def test_import_des_utilisateurs_ne_chiffre_pas_pendant_la_verification(
    admin_magasin, tunis, monkeypatch
):
    # Chiffrer un mot de passe prend près d'une seconde sur un petit serveur : la vérification
    # (annulée) ne chiffre rien, l'import chiffre une fois chaque mot de passe.
    from django.contrib.auth import base_user

    from apps.securite import imports

    appels = []
    chiffrer = imports.make_password

    def un_a_un(mot, *args):
        # make_password(None) donne un mot de passe inutilisable, sans calcul.
        return chiffrer(mot) if mot is None else pytest.fail("chiffré un à un")

    monkeypatch.setattr(base_user, "make_password", un_a_un)
    monkeypatch.setattr(imports, "make_password", lambda m: appels.append(m) or chiffrer(m))
    demandeur = Utilisateur.objects.get(username="admin")
    lignes = [
        (
            2,
            {
                "identifiant": n,
                "mot_de_passe": "Provisoire-2026!",
                "profil": "Vendeur",
                "magasin": "T01",
                "prenom": n,
                "nom": "Test",
            },
        )
        for n in ("sarra", "olfa")
    ]
    assert imports.importer_utilisateurs(lignes, demandeur=demandeur, apercu=True).crees == 2
    assert appels == [] and not Utilisateur.objects.filter(username="sarra").exists()
    assert imports.importer_utilisateurs(lignes, demandeur=demandeur).crees == 2
    assert appels == ["Provisoire-2026!"]
    assert Utilisateur.objects.get(username="olfa").check_password("Provisoire-2026!")


def test_import_des_utilisateurs_garde_les_garde_fous(admin_magasin, reseau):
    # Un administrateur de magasin ne donne pas de profil sur tout le réseau.
    contenu = fichier(
        "identifiant;mot_de_passe;profil\nkarim;Provisoire-2026!;Vendeur\nnadia;court;Vendeur\n"
    )
    erreurs = importer(admin_magasin, contenu, "utilisateurs").json()["erreurs"]
    assert [e["ligne"] for e in erreurs] == [2, 3]
    assert "Provisoire" not in str(erreurs) and "court" not in str(erreurs)
    assert not Utilisateur.objects.filter(username__in=["karim", "nadia"]).exists()
