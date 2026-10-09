"""Imports Excel / CSV depuis l'administration du serveur (/admin/)."""

import io

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from apps.achats.models import Fournisseur

FICHIER = "raison_sociale;matricule_fiscal;ville\nHoya Lens Tunisie;7777777/H/M/000;Sfax\n"


def _admin(creer_utilisateur, client_de):
    anas = creer_utilisateur("anas")
    anas.is_staff = anas.is_superuser = True
    anas.save()
    return client_de(anas)


def test_import_des_fournisseurs_depuis_l_administration(creer_utilisateur, client_de, tunis):
    navigateur = _admin(creer_utilisateur, client_de)
    liste = navigateur.get("/admin/achats/fournisseur/").content.decode()
    assert 'href="/admin/achats/fournisseur/importer/fournisseurs/"' in liste

    modele = navigateur.get("/admin/achats/fournisseur/importer/fournisseurs/modele.xlsx")
    assert load_workbook(io.BytesIO(modele.content)).sheetnames == ["À remplir", "Exemple", "Aide"]

    url = "/admin/achats/fournisseur/importer/fournisseurs/"
    page = navigateur.get(url)
    assert page.status_code == 200 and str(tunis.public_id) in page.content.decode()

    fichier = SimpleUploadedFile("fournisseurs.csv", FICHIER.encode())
    verification = navigateur.post(url, {"fichier": fichier, "magasin": str(tunis.public_id)})
    assert verification.context["rapport"].crees == 1
    assert not Fournisseur.objects.exists()
    jeton = verification.context["jeton"]

    # Le fichier vérifié est gardé : l'import ne le renvoie pas, et un autre jeton est refusé.
    refus = navigateur.post(
        url, {"importer": "1", "jeton": "faux", "magasin": str(tunis.public_id)}
    )
    assert refus.status_code == 302 and not Fournisseur.objects.exists()
    fin = navigateur.post(url, {"importer": "1", "jeton": jeton, "magasin": str(tunis.public_id)})
    assert "Import terminé : 1 créé(s)" in fin.content.decode()
    hoya = Fournisseur.objects.get()
    assert (hoya.nom, hoya.ville, hoya.pays) == ("Hoya Lens Tunisie", "Sfax", tunis.pays)


def test_import_dans_l_administration_exige_les_droits(affecter, client_de, tunis):
    lecteur = affecter("lecteur", "achats.view_fournisseur", portee="reseau")
    lecteur.is_staff = True
    lecteur.save()
    navigateur = client_de(lecteur)
    liste = navigateur.get("/admin/achats/fournisseur/").content.decode()
    assert "importer/fournisseurs" not in liste
    assert navigateur.get("/admin/achats/fournisseur/importer/fournisseurs/").status_code == 403
    assert navigateur.get("/admin/achats/fournisseur/importer/verres/").status_code == 404


def test_boutons_d_import_des_listes(creer_utilisateur, client_de):
    navigateur = _admin(creer_utilisateur, client_de)
    for liste, imports in (
        ("/admin/stock/article/", ["catalogue", "verres"]),
        ("/admin/stock/mouvementstock/", ["stock"]),
        ("/admin/achats/bonreception/", ["receptions"]),
        ("/admin/crm/client/", ["clients"]),
        ("/admin/securite/utilisateur/", ["utilisateurs"]),
    ):
        page = navigateur.get(liste).content.decode()
        for type_import in imports:
            assert f"{liste}importer/{type_import}/" in page
