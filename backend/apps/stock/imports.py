"""Import du catalogue et des entrées de stock depuis un fichier Excel (.xlsx) ou CSV.

Une ligne d'en-tête nomme les colonnes (casse, accents et espaces sans importance). Le fichier
est contrôlé en entier : à la moindre erreur, rien n'est enregistré et chaque ligne fautive est
signalée avec son numéro. ``apercu=True`` fait le même contrôle sans rien enregistrer ; il
signale aussi en alerte les articles déjà au catalogue ou déjà en stock. L'import lui-même
exige le jeton de cette vérification, propre au fichier : on ne peut pas importer un fichier
qu'on n'a pas vérifié.
"""

import csv
import hashlib
import hmac
import io
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Sum

from apps.achats.models import Fournisseur
from apps.reseau.models import TauxTva

from .models import (
    Article,
    Lentille,
    Monture,
    MouvementStock,
    PlageVerre,
    PrixArticle,
    Verre,
    stock_disponible,
)

TAILLE_MAX = 5 * 1024 * 1024
LIGNES_MAX = 20000


class FichierIllisible(Exception):
    pass


@dataclass
class Rapport:
    apercu: bool
    lignes: int = 0
    crees: int = 0
    modifies: int = 0
    erreurs: list = field(default_factory=list)
    alertes: list = field(default_factory=list)
    jeton: str = ""

    def erreur(self, ligne, message):
        self.erreurs.append({"ligne": ligne, "message": message})

    def alerte(self, ligne, message):
        self.alertes.append({"ligne": ligne, "message": message})


def jeton_de_verification(contenu, *contexte):
    """Signature du fichier vérifié (et de son magasin) : l'import exige la même."""
    message = hashlib.sha256(contenu).hexdigest() + "|" + "|".join(str(c) for c in contexte)
    return hmac.new(settings.SECRET_KEY.encode(), message.encode(), hashlib.sha256).hexdigest()


def _stocks_par_magasin(article):
    """« Tunis Centre 4, Sfax 2 » : stock de l'article dans les magasins du périmètre."""
    lignes = (
        MouvementStock.objects.filter(article=article)
        .values("magasin__nom")
        .annotate(total=Sum("quantite"))
        .order_by("magasin__nom")
    )
    return ", ".join(
        f"{ligne['magasin__nom']} {ligne['total']}" for ligne in lignes if ligne["total"]
    )


class _Annuler(Exception):
    """Annule la transaction d'un import en erreur ou d'un aperçu."""


def normaliser(texte):
    texte = unicodedata.normalize("NFKD", str(texte or "")).encode("ascii", "ignore").decode()
    return "_".join(texte.strip().lower().replace("-", " ").replace("'", " ").split())


def lire_tableau(fichier):
    """Lignes du fichier en dictionnaires {colonne normalisée: texte}, avec leur n° de ligne."""
    if fichier.size > TAILLE_MAX:
        raise FichierIllisible("Fichier trop volumineux (5 Mo au plus).")
    nom = (fichier.name or "").lower()
    contenu = fichier.read()
    if nom.endswith(".xlsx"):
        lignes = _lire_xlsx(contenu)
    elif nom.endswith((".csv", ".txt")):
        lignes = _lire_csv(contenu)
    else:
        raise FichierIllisible("Format non reconnu : envoyer un fichier .xlsx ou .csv.")
    lignes = iter(lignes)
    entetes = [normaliser(c) for c in next(lignes, [])]
    if not any(entetes):
        raise FichierIllisible("Le fichier est vide : la première ligne doit nommer les colonnes.")
    resultat = []
    for numero, valeurs in enumerate(lignes, start=2):
        texte = ["" if v is None else str(v).strip() for v in valeurs]
        if not any(texte):
            continue
        if len(resultat) >= LIGNES_MAX:
            raise FichierIllisible(f"Trop de lignes ({LIGNES_MAX} au plus par fichier).")
        resultat.append((numero, dict(zip(entetes, texte, strict=False))))
    return resultat


def _lire_xlsx(contenu):
    from openpyxl import load_workbook

    try:
        classeur = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    except Exception as erreur:  # fichier corrompu ou d'un autre format
        raise FichierIllisible("Fichier Excel illisible.") from erreur
    feuille = classeur.worksheets[0]
    for valeurs in feuille.iter_rows(values_only=True):
        yield [_texte_cellule(v) for v in valeurs]


def _texte_cellule(valeur):
    # Excel rend 13 chiffres d'un code-barres en float (8.05e12) : on les remet en entier.
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    return valeur


def _lire_csv(contenu):
    for codage in ("utf-8-sig", "cp1252"):
        try:
            texte = contenu.decode(codage)
            break
        except UnicodeDecodeError:
            continue
    premiere = texte.split("\n", 1)[0]
    separateur = ";" if premiere.count(";") >= premiere.count(",") else ","
    return list(csv.reader(io.StringIO(texte), delimiter=separateur))


def _decimal(texte, nom):
    if texte == "":
        return None
    try:
        return Decimal(texte.replace(" ", "").replace(",", "."))
    except InvalidOperation:
        raise ValidationError(f"{nom} : « {texte} » n'est pas un nombre.") from None


def _entier(texte, nom):
    valeur = _decimal(texte, nom)
    if valeur is None:
        return None
    if valeur != valeur.to_integral_value():
        raise ValidationError(f"{nom} : « {texte} » n'est pas un nombre entier.")
    return int(valeur)


def _booleen(texte):
    return normaliser(texte) in {"oui", "o", "x", "1", "vrai", "true", "yes"}


def _choix(texte, choix, nom):
    if texte == "":
        return ""
    cle = normaliser(texte)
    for valeur, libelle in choix:
        if cle in (normaliser(valeur), normaliser(libelle)):
            return valeur
    permis = ", ".join(str(libelle) for _, libelle in choix)
    raise ValidationError(f"{nom} : « {texte} » inconnu (valeurs possibles : {permis}).")


FICHES = {
    Article.Famille.MONTURE: Monture,
    Article.Famille.VERRE: Verre,
    Article.Famille.LENTILLE: Lentille,
}

# Plage de puissances d'un verre : une ligne par plage, même référence répétée.
COLONNES_PLAGE = ["sphere_debut", "sphere_fin", "cylindre_debut", "cylindre_fin"]
# Colonnes du modèle d'import, dans l'ordre du fichier modèle.
COLONNES_CATALOGUE = [
    "reference",
    "libelle",
    "famille",
    "fournisseur",
    "reference_fournisseur",
    "code_barres",
    "sur_commande",
    "prix_ttc",
    "tva",
    "prix_achat_ht",
    "categorie",
    "marque",
    "modele",
    "couleur",
    "couleur_verres",
    "matiere",
    "type",
    "forme",
    "genre",
    "tranche_age",
    "calibre",
    "pont",
    "branche",
    "gamme",
    "geometrie",
    "indice",
    "traitements",
    "photochromique",
    "teinte",
    "diametre",
    "fabrication",
    "diametre_commercial",
    *COLONNES_PLAGE,
    "renouvellement",
    "rayon",
    "puissance",
    "cylindre",
    "axe",
    "addition",
    "lentilles_par_boite",
]
COLONNES_STOCK = ["code_barres", "reference", "quantite"]
# Modèle « verres » : les colonnes du catalogue utiles aux verres, famille implicite.
COLONNES_VERRES = [
    "reference",
    "libelle",
    "fournisseur",
    "reference_fournisseur",
    "code_barres",
    "sur_commande",
    "prix_ttc",
    "tva",
    "gamme",
    "geometrie",
    "indice",
    "matiere",
    "traitements",
    "photochromique",
    "teinte",
    "diametre",
    "fabrication",
    "diametre_commercial",
    *COLONNES_PLAGE,
    "prix_achat_ht",
]


def _valeurs_fiche(modele, ligne):
    """Champs de la fiche présents dans la ligne, convertis selon leur type."""
    valeurs = {}
    for champ in modele._meta.concrete_fields:
        nom = champ.name
        if nom == "article" or nom not in ligne:
            continue
        texte = ligne[nom]
        if texte == "" and not champ.blank:
            continue  # Case vide : valeur par défaut, ou valeur déjà enregistrée.
        if champ.choices:
            valeurs[nom] = _choix(texte, champ.choices, nom)
        elif champ.get_internal_type() == "BooleanField":
            valeurs[nom] = _booleen(texte)
        elif champ.get_internal_type() == "DecimalField":
            valeurs[nom] = _decimal(texte, nom)
        elif champ.get_internal_type() in ("PositiveSmallIntegerField", "IntegerField"):
            valeurs[nom] = _entier(texte, nom)
        else:
            valeurs[nom] = texte
    return valeurs


def _message(erreur):
    if hasattr(erreur, "message_dict"):
        return " ; ".join(
            f"{'' if cle == '__all__' else cle + ' : '}{' '.join(messages)}"
            for cle, messages in erreur.message_dict.items()
        )
    return " ".join(erreur.messages)


def importer_catalogue(lignes, *, pays, apercu=False, famille=None):
    """Crée ou met à jour les articles (par référence), leur fiche et leur prix dans ``pays``.

    Avec ``famille`` (import des verres…), la colonne famille est facultative et toute autre
    famille est refusée.
    """
    rapport = Rapport(apercu=apercu, lignes=len(lignes))
    obligatoires = {"reference", "libelle", "fournisseur"} | (set() if famille else {"famille"})
    manquantes = obligatoires - set(lignes[0][1] if lignes else {})
    if manquantes:
        rapport.erreur(1, f"Colonnes obligatoires absentes : {', '.join(sorted(manquantes))}.")
        return rapport
    fournisseurs = {normaliser(f.nom): f for f in Fournisseur.objects.all()} | {
        str(f.code): f for f in Fournisseur.objects.exclude(code=None)
    }
    taux = {t.taux: t for t in TauxTva.objects.filter(pays=pays)}
    vues, plages = {}, {}
    try:
        with transaction.atomic():
            for numero, ligne in lignes:
                reference = ligne.get("reference", "")
                avec_plage = any(ligne.get(nom, "") for nom in COLONNES_PLAGE)
                if reference in vues:
                    if not avec_plage:
                        rapport.erreur(
                            numero, f"Référence {reference} déjà en ligne {vues[reference]}."
                        )
                        continue
                    # Ligne suivante d'un même verre : une plage de puissances de plus.
                    article = Article.objects.filter(reference=reference).first()
                    if article is None:
                        rapport.erreur(
                            numero, f"{reference} : corriger d'abord la ligne {vues[reference]}."
                        )
                        continue
                    try:
                        with transaction.atomic():
                            _importer_plage(article, ligne, pays, plages)
                    except ValidationError as erreur:
                        rapport.erreur(numero, _message(erreur))
                    continue
                vues[reference] = numero
                existant = Article.objects.filter(reference=reference).first()
                try:
                    with transaction.atomic():
                        cree = _importer_article(ligne, pays, fournisseurs, taux, famille)
                        if avec_plage:
                            article = Article.objects.get(reference=reference)
                            _importer_plage(article, ligne, pays, plages)
                except ValidationError as erreur:
                    rapport.erreur(numero, _message(erreur))
                    continue
                if cree:
                    rapport.crees += 1
                else:
                    rapport.modifies += 1
                    en_stock = _stocks_par_magasin(existant)
                    rapport.alerte(
                        numero,
                        f"{reference} existe déjà au catalogue : il sera mis à jour"
                        + (f" (en stock : {en_stock})." if en_stock else " (pas en stock)."),
                    )
            if rapport.erreurs or apercu:
                raise _Annuler
    except _Annuler:
        pass
    return rapport


def _importer_article(ligne, pays, fournisseurs, taux, imposee=None):
    famille = _choix(ligne.get("famille", ""), Article.Famille.choices, "famille") or imposee
    if not famille:
        raise ValidationError("famille : obligatoire.")
    if imposee and famille != imposee:
        raise ValidationError(f"famille : ce fichier n'importe que des {imposee}s.")
    fournisseur = fournisseurs.get(normaliser(ligne["fournisseur"])) or fournisseurs.get(
        ligne["fournisseur"].strip()
    )
    if fournisseur is None:
        raise ValidationError(
            f"fournisseur : « {ligne['fournisseur']} » inconnu (code ou raison sociale) ; le "
            "créer d'abord (écran Fournisseurs ou import des fournisseurs)."
        )
    code_barres = ligne.get("code_barres", "")
    if code_barres:
        autre = Article.objects.filter(code_barres=code_barres).exclude(
            reference=ligne["reference"]
        )
        if autre.exists():
            raise ValidationError(
                f"code_barres : {code_barres} est déjà celui de {autre.first().reference}."
            )
    article = Article.objects.filter(reference=ligne["reference"]).first()
    proche = Article.objects.filter(reference__iexact=ligne["reference"]).first()
    if article is None and proche is not None:
        raise ValidationError(
            f"reference : {ligne['reference']} est déjà celle de {proche.reference} "
            "(majuscules et minuscules comptent pour la même référence)."
        )
    cree = article is None
    if cree:
        article = Article(reference=ligne["reference"])
    article.libelle = ligne["libelle"]
    article.famille = famille
    article.fournisseur = fournisseur
    for champ in ("reference_fournisseur", "code_barres"):
        if champ in ligne:
            setattr(article, champ, ligne[champ])
    if "sur_commande" in ligne:
        article.sur_commande = _booleen(ligne["sur_commande"])
    elif cree:
        article.sur_commande = famille == Article.Famille.VERRE
    article.full_clean()
    article.save()

    modele = FICHES.get(famille)
    if modele is not None:
        fiche = modele.objects.filter(article=article).first() or modele(article=article)
        valeurs = _valeurs_fiche(modele, ligne)
        # Anciens fichiers : colonne « solaire » (oui/non) au lieu de la famille de monture.
        if modele is Monture and "categorie" not in ligne and _booleen(ligne.get("solaire", "")):
            valeurs["categorie"] = Monture.Categorie.SOLAIRE
        for nom, valeur in valeurs.items():
            setattr(fiche, nom, valeur)
        fiche.full_clean()
        fiche.save()

    prix = _decimal(ligne.get("prix_ttc", ""), "prix_ttc")
    if prix is not None:
        taux_tva = _decimal(ligne.get("tva", ""), "tva")
        if taux_tva is None or taux_tva not in taux:
            permis = ", ".join(f"{t:g}" for t in sorted(taux))
            raise ValidationError(f"tva : taux du pays à préciser ({permis}).")
        tarif = PrixArticle.objects.filter(article=article, pays=pays).first() or PrixArticle(
            article=article, pays=pays
        )
        tarif.prix_vente_ttc, tarif.tva = prix, taux[taux_tva]
        achat = _decimal(ligne.get("prix_achat_ht", ""), "prix_achat_ht")
        if achat is not None:
            tarif.prix_achat_ht = achat
        tarif.full_clean()
        tarif.save()
    return cree


def _importer_plage(article, ligne, pays, plages):
    """Ajoute une plage de puissances au verre ; les plages d'avant (même pays) sont remplacées
    par celles du fichier. ``plages`` compte les plages déjà lues de chaque verre."""
    if article.famille != Article.Famille.VERRE:
        raise ValidationError("sphere_debut : les plages de puissances ne vont qu'avec un verre.")
    if article.pk not in plages:
        PlageVerre.objects.filter(article=article, pays=pays).delete()
        plages[article.pk] = 0
    valeurs = {nom: _decimal(ligne.get(nom, ""), nom) for nom in COLONNES_PLAGE}
    if valeurs["sphere_debut"] is None or valeurs["sphere_fin"] is None:
        raise ValidationError("sphere_debut et sphere_fin : obligatoires pour une plage.")
    prix = _decimal(ligne.get("prix_ttc", ""), "prix_ttc")
    if prix is None:
        raise ValidationError("prix_ttc : obligatoire pour une plage de puissances.")
    plages[article.pk] += 1
    plage = PlageVerre(
        article=article,
        pays=pays,
        ordre=plages[article.pk],
        sphere_debut=valeurs["sphere_debut"],
        sphere_fin=valeurs["sphere_fin"],
        cylindre_debut=valeurs["cylindre_debut"] or 0,
        cylindre_fin=valeurs["cylindre_fin"] or 0,
        prix_achat_ht=_decimal(ligne.get("prix_achat_ht", ""), "prix_achat_ht"),
        prix_vente_ttc=prix,
    )
    plage.full_clean()
    plage.save()


def importer_stock(lignes, *, magasin, utilisateur, piece="", apercu=False):
    """Réceptions de stock dans ``magasin`` : article par code-barres ou référence, quantité."""
    rapport = Rapport(apercu=apercu, lignes=len(lignes))
    colonnes = set(lignes[0][1]) if lignes else set()
    if "quantite" not in colonnes or not colonnes & {"code_barres", "reference"}:
        rapport.erreur(1, "Colonnes attendues : code_barres ou reference, et quantite.")
        return rapport
    stocks = {}
    try:
        with transaction.atomic():
            for numero, ligne in lignes:
                code, reference = ligne.get("code_barres", ""), ligne.get("reference", "")
                article = None
                if code:
                    article = Article.objects.filter(code_barres=code, est_actif=True).first()
                if article is None and reference:
                    article = Article.objects.filter(reference=reference, est_actif=True).first()
                if article is None:
                    rapport.erreur(numero, f"Article inconnu : {code or reference or '(vide)'}.")
                    continue
                if article.sur_commande:
                    rapport.erreur(
                        numero,
                        f"{article.reference} est commandé pour chaque client : pas de stock.",
                    )
                    continue
                try:
                    quantite = _entier(ligne["quantite"], "quantite")
                except ValidationError as erreur:
                    rapport.erreur(numero, _message(erreur))
                    continue
                if not quantite or quantite < 0:
                    rapport.erreur(numero, "quantite : un nombre entier positif.")
                    continue
                if article.pk not in stocks:
                    stocks[article.pk] = stock_disponible(magasin, article)
                avant = stocks[article.pk]
                stocks[article.pk] = avant + quantite
                if avant > 0:
                    rapport.alerte(
                        numero,
                        f"{article.reference} déjà en stock à {magasin.nom} : {avant} ; "
                        f"{avant + quantite} après l'entrée.",
                    )
                MouvementStock.tous.create(
                    magasin=magasin,
                    article=article,
                    quantite=quantite,
                    type=MouvementStock.Type.RECEPTION,
                    utilisateur=utilisateur,
                    reference=piece[:60],
                )
                rapport.crees += 1
            if rapport.erreurs or apercu:
                raise _Annuler
    except _Annuler:
        pass
    return rapport
