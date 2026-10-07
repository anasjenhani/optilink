"""Fichiers modèles des imports, à télécharger en Excel (.xlsx) ou en CSV.

Le modèle Excel a trois feuilles : « À remplir » (la ligne d'en-tête, seule feuille lue à
l'import), « Exemple » (quelques lignes remplies) et « Aide » (le sens de chaque colonne).
Le CSV n'a que la ligne d'en-tête, séparée par « ; » comme l'Excel français.
"""

import io
from dataclasses import dataclass

from apps.achats.imports import COLONNES_FOURNISSEURS, COLONNES_RECEPTIONS
from apps.crm.imports import COLONNES_CLIENTS
from apps.optique.imports import COLONNES_OPHTALMOLOGUES
from apps.securite.imports import COLONNES_UTILISATEURS
from apps.stock.imports import COLONNES_CATALOGUE, COLONNES_STOCK, COLONNES_VERRES


@dataclass(frozen=True)
class Modele:
    titre: str
    colonnes: list
    obligatoires: set
    aide: dict
    exemples: list


_COMMUN_ARTICLES = {
    "reference": "Référence interne ; réimporter la même référence met l'article à jour.",
    "libelle": "Désignation affichée en vente.",
    "fournisseur": "Code ou raison sociale d'un fournisseur déjà créé.",
    "reference_fournisseur": "Référence de l'article chez le fournisseur.",
    "code_barres": "Code-barres (EAN) ; il ne peut appartenir qu'à un seul article.",
    "sur_commande": "oui si l'article est commandé pour chaque client (verres), sinon non.",
    "prix_ttc": "Prix de vente TTC.",
    "tva": "Taux de TVA du prix de vente (7, 13, 19…).",
}

_AIDE_PLAGES = {
    "fabrication": "Verre sur commande : « stock » (stock fournisseur) ou « prescription » "
    "(RX, importation, par défaut).",
    "diametre_commercial": "Diamètre commercial (65/70, 70/75…).",
    "sphere_debut": "Plage de puissances : sphère de début. Une ligne par plage, en répétant "
    "la référence ; prix_ttc est alors le prix de la plage. Les plages du fichier remplacent "
    "celles déjà enregistrées pour ce verre.",
    "sphere_fin": "Plage de puissances : sphère de fin.",
    "cylindre_debut": "Plage de puissances : cylindre de début (0 si vide).",
    "cylindre_fin": "Plage de puissances : cylindre de fin (0 si vide).",
    "prix_achat_ht": "Prix d'achat HT avant remise (de la plage s'il y en a une).",
}

MODELES = {
    "verres": Modele(
        titre="Verres",
        colonnes=COLONNES_VERRES,
        obligatoires={"reference", "libelle", "fournisseur", "geometrie"},
        aide=_COMMUN_ARTICLES
        | {
            "gamme": "Nom commercial du verre (Varilux Comfort…).",
            "geometrie": "Unifocal, Progressif, Dégressif ou Bifocal.",
            "indice": "Indice de réfraction (1.5, 1.6, 1.67…).",
            "matiere": "Organique, Polycarbonate ou Minéral.",
            "traitements": "Antireflet, durci, filtre lumière bleue…",
            "photochromique": "oui ou non.",
            "teinte": "Teinte éventuelle.",
            "diametre": "Diamètre en mm.",
        }
        | _AIDE_PLAGES,
        exemples=[
            {
                "reference": "VER-ST-156",
                "libelle": "RELAX 400 ASP 1.56 BLANC",
                "fournisseur": "TN OPTIC",
                "sur_commande": "oui",
                "prix_ttc": "92,859",
                "tva": "19",
                "geometrie": "Unifocal",
                "indice": "1,56",
                "fabrication": "stock",
                "diametre_commercial": "65",
                "sphere_debut": "-2,75",
                "sphere_fin": "4",
                "cylindre_debut": "0",
                "cylindre_fin": "3",
            },
            {
                "reference": "VER-ST-156",
                "libelle": "RELAX 400 ASP 1.56 BLANC",
                "fournisseur": "TN OPTIC",
                "prix_ttc": "98,573",
                "tva": "19",
                "geometrie": "Unifocal",
                "sphere_debut": "-4",
                "sphere_fin": "0",
                "cylindre_debut": "0",
                "cylindre_fin": "3",
            },
            {
                "reference": "VER-PR-016",
                "libelle": "Verre progressif 1.6 antireflet",
                "fournisseur": "Essilor Tunisie",
                "reference_fournisseur": "VX-COMF-16",
                "sur_commande": "oui",
                "prix_ttc": "320,000",
                "tva": "7",
                "gamme": "Varilux Comfort",
                "geometrie": "Progressif",
                "indice": "1,6",
                "matiere": "Organique",
                "traitements": "Antireflet",
                "photochromique": "non",
                "diametre": "70",
            },
        ],
    ),
    "catalogue": Modele(
        titre="Catalogue",
        colonnes=COLONNES_CATALOGUE,
        obligatoires={"reference", "libelle", "famille", "fournisseur"},
        aide=_COMMUN_ARTICLES
        | {
            "famille": "Monture, Verre, Lentille, Divers ou Supplément verre.",
            "prix_achat_ht": "Prix d'achat HT avant remise (avec prix_ttc).",
            "categorie": "Montures : Lunette Optique (par défaut), Lunette Solaire ou Lunette "
            "Applique.",
            "matiere": "Montures : Acétate, Titane, Acier, TR90, Corne, Bois ou Métal. "
            "Verres : Organique, Polycarbonate ou Minéral.",
            "type": "Montures : Cerclée, Semi-cerclée ou Percée (vide = aucun).",
            "forme": "Ronde, rectangle, papillon…",
            "tranche_age": "Adulte, Junior, Enfant ou Bébé.",
            "calibre": "Taille (mm).",
        }
        | _AIDE_PLAGES,
        exemples=[
            {
                "reference": "MON-RB5154",
                "libelle": "Ray-Ban RB5154 écaille",
                "famille": "Monture",
                "fournisseur": "Luxottica Tunisie",
                "code_barres": "8053672000001",
                "prix_ttc": "450,000",
                "tva": "19",
                "prix_achat_ht": "220,000",
                "categorie": "Lunette Optique",
                "marque": "Ray-Ban",
                "modele": "RB5154",
                "couleur": "Écaille",
                "matiere": "Acétate",
                "type": "Cerclée",
                "forme": "Clubmaster",
                "tranche_age": "Adulte",
                "calibre": "51",
                "pont": "21",
                "branche": "145",
            }
        ],
    ),
    "stock": Modele(
        titre="Entrées de stock",
        colonnes=COLONNES_STOCK,
        obligatoires={"quantite"},
        aide={
            "code_barres": "Code-barres de l'article (ou la référence).",
            "reference": "Référence de l'article si pas de code-barres.",
            "quantite": "Nombre d'articles reçus.",
        },
        exemples=[{"code_barres": "8053672000001", "quantite": "3"}],
    ),
    "clients": Modele(
        titre="Clients",
        colonnes=COLONNES_CLIENTS,
        obligatoires={"nom", "prenom"},
        aide={
            "ancien_numero": "N° de fiche dans l'ancien logiciel ; réimporter met la fiche à jour.",
            "civilite": "M ou Mme.",
            "date_naissance": "jj/mm/aaaa.",
            "accepte_relances": "oui ou non.",
        },
        exemples=[
            {
                "ancien_numero": "1024",
                "civilite": "Mme",
                "nom": "BEN SALAH",
                "prenom": "Leila",
                "date_naissance": "12/05/1980",
                "telephone": "22123456",
                "ville": "Tunis",
                "accepte_relances": "oui",
            }
        ],
    ),
    "fournisseurs": Modele(
        titre="Fournisseurs",
        colonnes=COLONNES_FOURNISSEURS,
        obligatoires={"raison_sociale"},
        aide={
            "code": "Vide pour un nouveau fournisseur (code attribué) ; le code pour modifier. "
            "Pour un fournisseur existant, une case vide garde la valeur déjà enregistrée.",
            "raison_sociale": "Nom du fournisseur ; il ne peut exister qu'une fois.",
            "notre_code": "Notre code client chez ce fournisseur.",
            "fournisseur_verres": "oui s'il fournit des verres.",
            "matricule_fiscal": "Ne peut appartenir qu'à un seul fournisseur.",
            "forme_juridique": "SARL, SUARL, SA, SNC ou PP.",
            "timbre_fiscal": "oui ou non (oui par défaut).",
            "assujetti": "oui ou non (oui par défaut).",
            "fodec": "oui si ses factures portent le FODEC (1 %).",
            "regime_tva": "Payer TVA, Export ou Exonération.",
            "exoneration_du": "Début d'exonération, jj/mm/aaaa.",
            "exoneration_au": "Fin d'exonération, jj/mm/aaaa.",
            "pays": "Code du pays (TN par défaut).",
        },
        exemples=[
            {
                "raison_sociale": "Essilor Tunisie",
                "responsable": "M. Trabelsi",
                "fournisseur_verres": "oui",
                "matricule_fiscal": "1234567/A/M/000",
                "forme_juridique": "SARL",
                "fodec": "oui",
                "regime_tva": "Payer TVA",
                "adresse": "Zone industrielle Charguia",
                "ville": "Tunis",
                "pays": "TN",
                "telephone": "71000000",
            }
        ],
    ),
    "receptions": Modele(
        titre="Bons de réception",
        colonnes=COLONNES_RECEPTIONS,
        obligatoires={"numero_bl", "fournisseur", "quantite", "prix_achat_ht"},
        aide={
            "numero_bl": "N° du bon de livraison ; les lignes d'un même BL forment un seul bon.",
            "date_bl": "Date du BL, jj/mm/aaaa (aujourd'hui si vide).",
            "fournisseur": "Code ou raison sociale d'un fournisseur déjà créé.",
            "code_barres": "Code-barres de l'article reçu (ou la référence).",
            "reference": "Référence de l'article si pas de code-barres.",
            "quantite": "Nombre d'articles reçus.",
            "prix_achat_ht": "Prix d'achat unitaire HT.",
            "taux_remise": "Remise en % (0 si vide).",
            "taux_tva": "Taux de TVA (celui de l'article si vide).",
            "numero_lot": "Lentilles et produits.",
            "date_peremption": "jj/mm/aaaa.",
            "observation": "Remarque sur le bon.",
        },
        exemples=[
            {
                "numero_bl": "BL-2026-0458",
                "date_bl": "05/10/2026",
                "fournisseur": "Luxottica Tunisie",
                "code_barres": "8053672000001",
                "quantite": "2",
                "prix_achat_ht": "210,000",
                "taux_remise": "5",
                "taux_tva": "19",
            },
            {
                "numero_bl": "BL-2026-0458",
                "date_bl": "05/10/2026",
                "fournisseur": "Luxottica Tunisie",
                "reference": "MON-RB3025",
                "quantite": "1",
                "prix_achat_ht": "180,000",
            },
        ],
    ),
    "ophtalmologues": Modele(
        titre="Ophtalmologistes",
        colonnes=COLONNES_OPHTALMOLOGUES,
        obligatoires={"nom"},
        aide={
            "code": "Code du médecin dans OptiLink (001, 002…) ; il ne peut appartenir qu'à "
            "un seul médecin.",
            "ancien_code": "Code du médecin dans l'ancien logiciel (CodeMedecin), gardé pour "
            "reprendre ses anciennes ordonnances.",
            "nom": "Nom de famille (ou nom complet). Un médecin déjà dans la liste, même écrit "
            "autrement (Dr, accents, majuscules, nom et prénom inversés), n'est pas recréé : "
            "sa fiche est complétée.",
            "prenom": "Prénom.",
            "telephone": "Téléphone du cabinet.",
            "telephone_2": "Portable.",
            "ville": "Ville (Tunis, Ariana…).",
        },
        exemples=[
            {
                "code": "001",
                "ancien_code": "001",
                "nom": "Belhaj",
                "prenom": "Taieb",
                "telephone": "71860266",
                "telephone_2": "24344866",
                "adresse": "Clinique du Lac",
                "ville": "Tunis",
            }
        ],
    ),
    "utilisateurs": Modele(
        titre="Utilisateurs",
        colonnes=COLONNES_UTILISATEURS,
        obligatoires={"identifiant", "mot_de_passe", "profil"},
        aide={
            "identifiant": "Identifiant de connexion ; un compte existant n'est pas modifié.",
            "mot_de_passe": "Mot de passe provisoire (12 caractères au moins). Supprimer le "
            "fichier après l'import.",
            "profil": "Vendeur, Caissier, Opticien, Responsable de magasin…",
            "magasin": "Code du magasin (MAG001…) pour un profil limité à ce magasin.",
            "societe": "Code de la société pour un profil sur toute la société.",
            "actif": "oui (par défaut) ou non : un compte inactif ne peut pas se connecter.",
        },
        exemples=[
            {
                "identifiant": "sarra",
                "prenom": "Sarra",
                "nom": "Jlassi",
                "email": "sarra@exemple.tn",
                "mot_de_passe": "Provisoire-2026!",
                "profil": "Vendeur",
                "magasin": "MAG001",
            },
            {
                "identifiant": "sarra",
                "profil": "Caissier",
                "magasin": "MAG001",
                "mot_de_passe": "Provisoire-2026!",
            },
        ],
    ),
}


def fichier_csv(modele):
    return ("﻿" + ";".join(modele.colonnes) + "\n").encode("utf-8")


def fichier_xlsx(modele):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    classeur = Workbook()
    entete = Font(bold=True, color="FFFFFF")
    fond = PatternFill("solid", fgColor="8B1414")

    def feuille(titre, lignes, premiere=None):
        f = premiere or classeur.create_sheet(titre)
        f.title = titre
        for ligne in lignes:
            f.append(ligne)
        for cellule in f[1]:
            cellule.font, cellule.fill = entete, fond
        for colonne in f.columns:
            largeur = max(len(str(c.value or "")) for c in colonne)
            f.column_dimensions[colonne[0].column_letter].width = min(max(largeur + 2, 12), 60)
        f.freeze_panes = "A2"
        return f

    a_remplir = feuille("À remplir", [modele.colonnes], classeur.active)
    # Tout en texte : Excel ne transforme ni les codes-barres ni les n° de BL.
    for lettre in range(1, len(modele.colonnes) + 1):
        a_remplir.column_dimensions[a_remplir.cell(1, lettre).column_letter].number_format = "@"
    feuille(
        "Exemple",
        [modele.colonnes] + [[e.get(c, "") for c in modele.colonnes] for e in modele.exemples],
    )
    feuille(
        "Aide",
        [["Colonne", "Obligatoire", "Contenu"]]
        + [
            [c, "oui" if c in modele.obligatoires else "", modele.aide.get(c, "")]
            for c in modele.colonnes
        ],
    )
    sortie = io.BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()
